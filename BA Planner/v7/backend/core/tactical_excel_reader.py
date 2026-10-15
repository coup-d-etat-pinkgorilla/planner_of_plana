"""Bounded, write-free spreadsheet input for FB7; no battle interpretation."""
from __future__ import annotations

from dataclasses import dataclass
import csv
import hashlib
from io import BytesIO, StringIO
import math
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree.ElementTree import ParseError
import zipfile

import openpyxl
from defusedxml.common import DefusedXmlException


class SourceReadError(ValueError):
    def __init__(self, reason: str) -> None:
        self.reason = reason
        super().__init__(reason)


@dataclass(frozen=True)
class ReaderLimits:
    file_bytes: int = 16 * 1024 * 1024
    zip_members: int = 2048
    zip_uncompressed_bytes: int = 64 * 1024 * 1024
    sheets: int = 50
    rows: int = 10_000
    columns: int = 64
    cell_utf8_bytes: int = 4096


@dataclass(frozen=True)
class RawCell:
    row: int
    column: int
    value: Any
    data_type: str
    number_format: str | None = None
    formula: str | None = None
    cached_value: Any = None


@dataclass(frozen=True)
class SourceDocument:
    filename: str
    format: str
    fingerprint: str
    sheet_names: tuple[str, ...]
    selected_sheet: str
    rows: tuple[tuple[RawCell, ...], ...]


def _check_cell(value: Any, limits: ReaderLimits) -> None:
    if isinstance(value, str) and len(value.encode('utf-8')) > limits.cell_utf8_bytes:
        raise SourceReadError('cell_limit_exceeded')
    if isinstance(value, float) and not math.isfinite(value):
        raise SourceReadError('non_finite_cell')


def _source_bytes(path: Path, limits: ReaderLimits) -> bytes:
    try:
        with path.open('rb') as handle:
            data = handle.read(limits.file_bytes + 1)
    except OSError as exc:
        raise SourceReadError('source_unavailable') from exc
    if len(data) > limits.file_bytes:
        raise SourceReadError('file_limit_exceeded')
    return data


def _zip_preflight(data: bytes, limits: ReaderLimits) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > limits.zip_members:
                raise SourceReadError('zip_member_limit_exceeded')
            if sum(info.file_size for info in infos) > limits.zip_uncompressed_bytes:
                raise SourceReadError('zip_size_limit_exceeded')
            for info in infos:
                name = PurePosixPath(info.filename)
                if name.is_absolute() or '..' in name.parts or '\\' in info.filename:
                    raise SourceReadError('invalid_zip_path')
                if info.flag_bits & 1:
                    raise SourceReadError('encrypted_workbook_unsupported')
                if info.filename.lower().endswith('vbaproject.bin'):
                    raise SourceReadError('macro_workbook_unsupported')
    except zipfile.BadZipFile as exc:
        raise SourceReadError('invalid_workbook') from exc


def _read_xlsx(data: bytes, sheet_name: str | None, limits: ReaderLimits):
    if not openpyxl.xml.DEFUSEDXML:
        raise SourceReadError('xml_guard_missing')
    _zip_preflight(data, limits)
    formulas = cached = None
    try:
        formulas = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False,
                                         keep_links=False, keep_vba=False)
        cached = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True,
                                       keep_links=False, keep_vba=False)
        names = tuple(formulas.sheetnames)
        if not names or len(names) > limits.sheets:
            raise SourceReadError('sheet_limit_exceeded')
        if sheet_name is None:
            visible = [sheet.title for sheet in formulas.worksheets if sheet.sheet_state == 'visible']
            if not visible:
                raise SourceReadError('no_visible_sheet')
            selected = visible[0]
        elif sheet_name not in names:
            raise SourceReadError('sheet_not_found')
        else:
            selected = sheet_name
        sheet = formulas[selected]
        if (sheet.max_row or 0) > limits.rows:
            raise SourceReadError('row_limit_exceeded')
        if (sheet.max_column or 0) > limits.columns:
            raise SourceReadError('column_limit_exceeded')
        cached_rows = iter(cached[selected].iter_rows())
        rows = []
        for row_number, cells in enumerate(sheet.iter_rows(), 1):
            if row_number > limits.rows:
                raise SourceReadError('row_limit_exceeded')
            if len(cells) > limits.columns:
                raise SourceReadError('column_limit_exceeded')
            cache_cells = next(cached_rows, ())
            row = []
            for column, cell in enumerate(cells, 1):
                value = cell.value
                _check_cell(value, limits)
                formula = value if cell.data_type == 'f' else None
                cached_value = cache_cells[column - 1].value if column <= len(cache_cells) else None
                _check_cell(cached_value, limits)
                row.append(RawCell(row_number, column, value, cell.data_type,
                                   getattr(cell, 'number_format', None), formula, cached_value))
            rows.append(tuple(row))
        return names, selected, tuple(rows)
    except SourceReadError:
        raise
    except DefusedXmlException as exc:
        raise SourceReadError('unsafe_xml') from exc
    except (OSError, ValueError, KeyError, TypeError, IndexError, ParseError, zipfile.BadZipFile) as exc:
        # openpyxl may wrap the XML guard's ValueError; retain its reason.
        cause = exc
        visited = set()
        while cause is not None and id(cause) not in visited:
            visited.add(id(cause))
            if isinstance(cause, DefusedXmlException):
                raise SourceReadError('unsafe_xml') from exc
            cause = cause.__cause__ or cause.__context__
        raise SourceReadError('invalid_workbook') from exc
    finally:
        if formulas is not None:
            formulas.close()
        if cached is not None:
            cached.close()


def _read_csv(data: bytes, delimiter: str, sheet_name: str | None, limits: ReaderLimits):
    if delimiter not in (',', '\t'):
        raise SourceReadError('delimiter_unsupported')
    if sheet_name not in (None, 'CSV'):
        raise SourceReadError('sheet_not_found')
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise SourceReadError('csv_encoding_unsupported') from exc
    rows = []
    try:
        for row_number, values in enumerate(csv.reader(StringIO(text, newline=''), delimiter=delimiter, strict=True), 1):
            if row_number > limits.rows:
                raise SourceReadError('row_limit_exceeded')
            if len(values) > limits.columns:
                raise SourceReadError('column_limit_exceeded')
            row = []
            for column, value in enumerate(values, 1):
                _check_cell(value, limits)
                # CSV content remains text; an Excel-looking string is not executed.
                row.append(RawCell(row_number, column, value, 's'))
            rows.append(tuple(row))
    except csv.Error as exc:
        raise SourceReadError('invalid_csv') from exc
    return ('CSV',), 'CSV', tuple(rows)


def read_source(path: str | Path, *, sheet_name: str | None = None,
                delimiter: str = ',', limits: ReaderLimits = ReaderLimits()) -> SourceDocument:
    source = Path(path)
    kind = source.suffix.lower()
    if kind not in ('.xlsx', '.csv'):
        raise SourceReadError('format_unsupported')
    data = _source_bytes(source, limits)
    if kind == '.xlsx':
        names, selected, rows = _read_xlsx(data, sheet_name, limits)
    else:
        names, selected, rows = _read_csv(data, delimiter, sheet_name, limits)
    return SourceDocument(source.name, kind[1:], hashlib.sha256(data).hexdigest(), names, selected, rows)
