from copy import deepcopy
from dataclasses import replace
from datetime import date, datetime, time, timedelta, timezone
import csv
import hashlib
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl
from core.repository_store import RepositoryError
from core.tactical_excel_reader import RawCell, SourceDocument, read_source
from core.tactical_excel_interpreter import (
    SCALARS, SLOTS, catalog_context, interpret, raw_key, validate_mapping, validate_resolutions, validate_rules,
)


FIELDS = (*SCALARS, *SLOTS)
MAPPING = {'header_row': 1, 'columns': {field: i for i, field in enumerate(FIELDS, 1)}}
OWN = ('shiroko', 'akane', 'yuuka', 'mika', 'serina', 'hibiki')
OPPONENT = {'version': 2, 'identity_id': 'opponent-one', 'current_display_name': 'Rival', 'aliases': ['Rival']}


def values(**changes):
    row = {'date': '2026-10-14', 'kind': 'attack', 'opponent': 'opponent-one',
           'season': '34', 'result': 'win', 'attempts': '2', 'notes': 'keep original'}
    for side in ('own', 'opponent'):
        for field, sid in zip((f for f in SLOTS if f.startswith(side + '_')), OWN):
            row[field] = sid
    row.update(changes)
    return row


def document(*rows, fingerprint='a' * 64):
    cells = [tuple(RawCell(1, col, field, 's') for col, field in enumerate(FIELDS, 1))]
    for number, data in enumerate(rows, 2):
        cells.append(tuple(RawCell(number, col, data.get(field), 's') for col, field in enumerate(FIELDS, 1)))
    return SourceDocument('synthetic.csv', 'csv', fingerprint, ('CSV',), 'CSV', tuple(cells))


def resolution(field, raw, action='value', value=None, row=None):
    data = {'field': field, 'raw': raw_key(raw), 'action': action, 'value': value}
    if row is not None:
        data['row'] = row
    return data


def decision(row, action, target=None):
    return {'row': row, 'action': action, 'target_match_id': target}


class ExcelInterpretationTests(unittest.TestCase):
    def setUp(self):
        self.context = catalog_context(set(OWN), [OPPONENT], [])

    def preview(self, *rows, **options):
        return interpret(document(*rows), self.context, column_mapping=MAPPING,
                         resolutions=options.pop('resolutions', []), row_decisions=options.pop('row_decisions', []), **options)

    def test_mapping_candidates_never_auto_approve(self):
        result = interpret(document(values()), self.context, column_mapping=None, resolutions=[], row_decisions=[])
        self.assertEqual([], result['rows'])
        self.assertEqual('mapping_confirmation_required', result['issues'][0]['code'])
        self.assertTrue(result['mapping_candidates'][0]['requires_confirmation'])

    def test_date_precision_positive_attempts_order_and_read_only_inputs(self):
        source, context = document(values()), deepcopy(self.context)
        before = deepcopy(context)
        result = interpret(source, context, column_mapping=MAPPING, resolutions=[], row_decisions=[])
        row = result['rows'][0]
        self.assertEqual('eligible', row['status'])
        self.assertEqual(('2026-10-14T00:00:00+09:00', 'date', 2),
                         tuple(row['match'][k] for k in ('occurred_at', 'occurred_precision', 'attempt_count')))
        self.assertEqual(list(OWN[:4]), [s['student_id'] for s in row['match']['attack_deck']['strikers']])
        self.assertEqual(before, context)
        self.assertEqual(1, result['summary']['eligible_rows'])

    def test_defense_reverses_own_and_rival_decks(self):
        result = self.preview(values(kind='defense', opponent_striker_1='izuna'))
        match = result['rows'][0]['match']
        self.assertEqual('shiroko', match['defense_deck']['strikers'][0]['student_id'])
        self.assertEqual('izuna', match['attack_deck']['strikers'][0]['student_id'])

    def test_missing_date_is_issue_not_now(self):
        result = self.preview(values(date=''))
        self.assertEqual(1, result['summary']['unresolved_rows'])
        self.assertIsNone(result['rows'][0]['match'])

    def test_naive_timestamp_requires_timezone(self):
        result = self.preview(values(date='2026-10-14T08:35:00'))
        self.assertEqual('timezone_required', result['rows'][0]['issues'][0]['code'])

    def test_zoned_timestamp_preserves_instant(self):
        row = self.preview(values(date='2026-10-14T08:35:00+09:00'))['rows'][0]
        self.assertEqual(('instant', '2026-10-13T23:35:00+00:00'),
                         (row['match']['occurred_precision'], row['match']['occurred_at']))

    def test_missing_attempts_bool_zero_fraction_negative_never_default(self):
        for raw in (None, '', 0, False, -1, 1.5, '1.5', '0'):
            with self.subTest(raw=raw):
                row = self.preview(values(attempts=raw))['rows'][0]
                self.assertIsNone(row['match'])
                self.assertIn('positive_attempts_required', [i['code'] for i in row['issues']])

    def test_combined_result_only_suggests_and_needs_separate_explicit_attempts(self):
        result = self.preview(values(result='승(2트)', attempts=''))
        group = next(g for g in result['string_groups'] if g['field'] == 'result')
        self.assertEqual({'result': 'win', 'attempt_count': 2, 'requires_confirmation': True}, group['suggestions'][0])
        resolved = self.preview(values(result='승(2트)', attempts=''), resolutions=[
            resolution('result', '승(2트)', value='win'), resolution('attempts', '', value=2)])
        self.assertEqual(1, resolved['summary']['eligible_rows'])
        self.assertEqual(1, len(resolved['rows']))

    def test_fuzzy_alias_groups_count_locations_without_approval(self):
        result = self.preview(values(own_striker_1='무장리오'), values(own_striker_1='무장리오', date='2026-10-15'))
        group = next(g for g in result['string_groups'] if g['field'] == 'own_striker_1')
        self.assertEqual(2, group['count'])
        self.assertEqual([2, 3], [p['row'] for p in group['places']])
        self.assertTrue(all(s['requires_confirmation'] for s in group['suggestions']))
        self.assertEqual(0, result['summary']['eligible_rows'])

    def test_canonical_unique_exact_name_resolves(self):
        name = self.context['students']['shiroko']['name']
        self.assertEqual('shiroko', self.preview(values(own_striker_1=name))['rows'][0]['match']['attack_deck']['strikers'][0]['student_id'])

    def test_form_name_ambiguity_does_not_choose_an_id(self):
        self.context['students']['akane']['name'] = self.context['students']['shiroko']['name']
        row = self.preview(values(own_striker_1=self.context['students']['shiroko']['name']))['rows'][0]
        self.assertIn('student_ambiguous', [i['code'] for i in row['issues']])

    def test_role_invalid_or_duplicate_or_not_owned_decks_are_unresolved(self):
        for row in (values(own_striker_1='serina'), values(own_striker_1='akane'), values(own_striker_1='izuna')):
            with self.subTest(row=row):
                self.assertEqual(0, self.preview(row)['summary']['eligible_rows'])

    def test_exact_opponent_display_name_is_only_suggestion(self):
        row = self.preview(values(opponent='Rival'))['rows'][0]
        self.assertIsNone(row['match'])
        accepted = self.preview(values(opponent='Rival'), resolutions=[resolution('opponent', 'Rival', value='opponent-one')])
        self.assertEqual('opponent-one', accepted['rows'][0]['match']['opponent_identity_id'])

    def test_new_opponent_requires_explicit_identity_and_alias_cannot_be_invented(self):
        resolved = self.preview(values(opponent='Other'), resolutions=[resolution('opponent', 'Other', value={'identity_id': 'new-other', 'display_name': 'Other'})])
        self.assertEqual('new-other', resolved['rows'][0]['match']['opponent_identity_id'])
        bad = self.preview(values(opponent='Other'), resolutions=[resolution('opponent', 'Other', value={'identity_id': 'opponent-one', 'display_name': 'Other'})])
        self.assertEqual('opponent_alias_required', bad['rows'][0]['issues'][0]['code'])

    def test_blank_ignore_and_wildcard_remain_distinct(self):
        for action, wildcard in (('ignore', False), ('wildcard', True)):
            result = self.preview(values(opponent_striker_1='?'), resolutions=[resolution('opponent_striker_1', '?', action=action)])
            slot = result['rows'][0]['match']['defense_deck']['strikers'][0]
            self.assertEqual((None, 'unknown', wildcard), (slot['student_id'], slot['state'], slot['wildcard']))
        slot = self.preview(values(opponent_striker_1=''))['rows'][0]['match']['defense_deck']['strikers'][0]
        self.assertEqual((None, 'unknown', False), (slot['student_id'], slot['state'], slot['wildcard']))

    def test_ignore_required_attempt_never_becomes_one(self):
        row = self.preview(values(attempts='?'), resolutions=[resolution('attempts', '?', action='ignore')])['rows'][0]
        self.assertIsNone(row['match'])

    def test_move_to_notes_preserves_text_without_student_identity(self):
        result = self.preview(values(opponent_striker_1='some note'), resolutions=[resolution('opponent_striker_1', 'some note', action='notes')])
        self.assertEqual('keep original\nsome note', result['rows'][0]['match']['notes'])
        self.assertIsNone(result['rows'][0]['match']['defense_deck']['strikers'][0]['student_id'])

    def test_numeric_notes_are_issue_not_slot_or_attempt(self):
        self.assertEqual(0, self.preview(values(notes=300))['summary']['eligible_rows'])

    def test_formula_cache_does_not_auto_approve(self):
        source = document(values())
        cells = list(source.rows[1]); index = FIELDS.index('attempts')
        cells[index] = RawCell(2, index + 1, '=1+1', 'f', None, '=1+1', 2)
        source = replace(source, rows=(source.rows[0], tuple(cells)))
        result = interpret(source, self.context, column_mapping=MAPPING, resolutions=[], row_decisions=[])
        self.assertEqual('formula_acceptance_required', result['rows'][0]['issues'][0]['code'])
        accepted = interpret(source, self.context, column_mapping=MAPPING,
                             resolutions=[resolution('attempts', '=1+1', value=2)], row_decisions=[])
        self.assertEqual(2, accepted['rows'][0]['match']['attempt_count'])

    def test_row_skip_and_empty_rows_do_not_change_denominator(self):
        empty = {field: '' for field in FIELDS}
        result = self.preview(values(), empty, values(attempts=''), row_decisions=[decision(4, 'skip')])
        self.assertEqual((3, 1, 1, 1, 0), tuple(result['summary'][k] for k in ('read_rows', 'empty_rows', 'skipped_rows', 'eligible_rows', 'unresolved_rows')))

    def test_skip_resolution_does_not_leave_group_from_the_skipped_row(self):
        result = self.preview(values(result='unknown', notes='skip'), resolutions=[resolution('notes', 'skip', action='skip')])
        self.assertEqual([], result['string_groups'])
        self.assertEqual('skipped', result['rows'][0]['status'])

    def test_saved_rule_must_be_live_and_role_valid(self):
        rules = {'version': 1, 'revision': 2, 'column_mapping': MAPPING,
                 'resolutions': [resolution('own_striker_1', 'alias', value='serina')]}
        row = self.preview(values(own_striker_1='alias'), rules=rules)['rows'][0]
        self.assertEqual('student_role_mismatch', row['issues'][0]['code'])
        self.assertTrue(row['issues'][0]['from_saved_rule'])
        explicit = self.preview(values(own_striker_1='alias'), rules=rules,
                                resolutions=[resolution('own_striker_1', 'alias', value='shiroko')])
        self.assertEqual(1, explicit['summary']['eligible_rows'])

    def test_saved_rules_do_not_auto_approve_column_mapping(self):
        rules = {'version': 1, 'revision': 1, 'column_mapping': MAPPING, 'resolutions': []}
        result = interpret(document(values()), self.context, column_mapping=None, resolutions=[], row_decisions=[], rules=rules)
        self.assertEqual(MAPPING, result['saved_mapping'])
        self.assertEqual(0, result['summary']['eligible_rows'])

    def test_resolution_must_bind_actual_raw_type_row_and_field(self):
        for res in (resolution('attempts', 2, value=3), resolution('result', 'win', value='loss', row=8)):
            with self.subTest(res=res), self.assertRaises(RepositoryError):
                self.preview(values(), resolutions=[res])

    def test_mapping_and_rule_shapes_reject_bool_duplicate_and_paths(self):
        bad = deepcopy(MAPPING); bad['columns']['date'] = True
        with self.assertRaises(RepositoryError): validate_mapping(bad)
        bad = deepcopy(MAPPING); bad['columns']['date'] = bad['columns']['kind']
        with self.assertRaises(RepositoryError): validate_mapping(bad)
        with self.assertRaises(RepositoryError): validate_rules({'version': 1, 'revision': 1, 'column_mapping': MAPPING, 'resolutions': [], 'source_path': 'private.xlsx'})
        with self.assertRaises(RepositoryError): validate_resolutions([resolution('date', '', action='wildcard')])

    def test_same_date_opponent_different_ordered_decks_are_separate(self):
        a, b = values(), values(own_striker_1='akane', own_striker_2='shiroko')
        result = self.preview(a, b)
        self.assertEqual((2, []), (result['summary']['eligible_rows'], result['conflicts']))
        self.assertNotEqual(result['rows'][0]['match']['match_id'], result['rows'][1]['match']['match_id'])

    def test_same_workbook_duplicate_needs_keep_separate_or_skip(self):
        result = self.preview(values(), values(notes='correction', attempts=3))
        self.assertEqual((0, 2), (result['summary']['eligible_rows'], result['summary']['conflict_rows']))
        accepted = self.preview(values(), values(notes='correction', attempts=3),
                                row_decisions=[decision(2, 'keep_separate'), decision(3, 'keep_separate')])
        self.assertEqual(2, accepted['summary']['eligible_rows'])
        self.assertEqual([2, 3], [r['match']['attempt_count'] for r in accepted['rows']])

    def test_coarse_existing_duplicate_overwrite_requires_explicit_target(self):
        item = self.preview(values())['rows'][0]['match']; item['match_id'] = 'old-match'
        self.context['matches'] = [item]
        result = self.preview(values(), on_conflict='overwrite')
        self.assertEqual('conflict', result['rows'][0]['status'])
        accepted = self.preview(values(), on_conflict='overwrite', row_decisions=[decision(2, 'overwrite', 'old-match')])
        self.assertEqual(('overwrite', 'old-match'), (accepted['rows'][0]['status'], accepted['rows'][0]['target_match_id']))

    def test_skip_conflict_preserves_old_and_explicit_keep_separate(self):
        item = self.preview(values())['rows'][0]['match']; item['match_id'] = 'old-match'
        self.context['matches'] = [item]
        result = self.preview(values())
        self.assertEqual('conflict_skipped', result['rows'][0]['status'])
        accepted = self.preview(values(), row_decisions=[decision(2, 'keep_separate')])
        self.assertEqual(1, accepted['summary']['eligible_rows'])
        self.assertEqual('old-match', self.context['matches'][0]['match_id'])

    def test_ambiguous_targets_and_double_overwrite_are_rejected(self):
        item = self.preview(values())['rows'][0]['match']; item['match_id'] = 'old-match'
        self.context['matches'] = [item]
        with self.assertRaises(RepositoryError):
            self.preview(values(), row_decisions=[decision(2, 'overwrite', 'unrelated')])
        with self.assertRaises(RepositoryError):
            self.preview(values(), values(), row_decisions=[decision(2, 'overwrite', 'old-match'), decision(3, 'overwrite', 'old-match')])

    def test_original_source_row_is_replay_not_new_match(self):
        row = self.preview(values())['rows'][0]
        item = row['match']; item['source_record_id'] = row['row_id']
        self.context['matches'] = [item]
        result = self.preview(values())
        self.assertEqual(('replay', 0), (result['rows'][0]['status'], result['summary']['eligible_rows']))

    def test_file_choices_rules_catalog_owned_and_identity_bind_interpretation_fp(self):
        before = self.preview(values())['interpretation_fingerprint']
        changes = [dict(on_conflict='overwrite'), dict(row_decisions=[decision(2, 'keep_separate')]),
                   dict(resolutions=[resolution('notes', 'keep original', value='edited')]),
                   dict(rules={'version': 1, 'revision': 1, 'column_mapping': MAPPING, 'resolutions': []})]
        for kwargs in changes:
            with self.subTest(kwargs=kwargs): self.assertNotEqual(before, self.preview(values(), **kwargs)['interpretation_fingerprint'])
        self.context['owned'].remove('mika')
        self.assertNotEqual(before, self.preview(values())['interpretation_fingerprint'])


class ActualExcelInputTests(unittest.TestCase):
    def test_time_only_and_duration_cells_are_serializable_review_issues(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'synthetic-time.xlsx'
            book = openpyxl.Workbook(); sheet = book.active
            sheet.append(list(FIELDS)); data = values()
            sheet.append([data[f] for f in FIELDS])
            sheet.cell(2, 1).value = time(12, 30); sheet.cell(2, 1).number_format = 'hh:mm'
            note_column = FIELDS.index('notes') + 1
            sheet.cell(2, note_column).value = timedelta(hours=1, minutes=30)
            sheet.cell(2, note_column).number_format = '[h]:mm:ss'
            book.save(path); book.close(); before = path.read_bytes()
            source = read_source(path)
            result = interpret(source, catalog_context(set(OWN), [OPPONENT], []), column_mapping=MAPPING,
                               resolutions=[], row_decisions=[])
            json.dumps(result, ensure_ascii=False, allow_nan=False)
            self.assertEqual(0, result['summary']['eligible_rows'])
            self.assertEqual({'date', 'notes'}, {i['field'] for i in result['rows'][0]['issues']})
            self.assertEqual('time', json.loads(result['rows'][0]['raw_cells'][0]['raw_key'])['type'])
            self.assertEqual('timedelta', json.loads(result['rows'][0]['raw_cells'][note_column - 1]['raw_key'])['type'])
            accepted = interpret(source, catalog_context(set(OWN), [OPPONENT], []), column_mapping=MAPPING,
                resolutions=[resolution('date', source.rows[1][0].value, value='2026-10-14'),
                             resolution('notes', source.rows[1][note_column-1].value, action='notes')], row_decisions=[])
            self.assertEqual(1, accepted['summary']['eligible_rows'])
            self.assertEqual('1:30:00', accepted['rows'][0]['match']['notes'])
            self.assertEqual(before, path.read_bytes())

    def test_actual_xlsx_date_formula_blank_and_source_preservation(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'synthetic.xlsx'
            book = openpyxl.Workbook(); sheet = book.active; sheet.title = 'Battles'
            sheet.append(list(FIELDS)); data = values()
            sheet.append([data[f] for f in FIELDS]); sheet.cell(2, 1).value = date(2026, 10, 14)
            sheet.cell(2, 1).number_format = 'yyyy-mm-dd'
            sheet.cell(2, 6).value = '=1+1'; book.save(path); book.close()
            before = path.read_bytes(); source = read_source(path)
            result = interpret(source, catalog_context(set(OWN), [OPPONENT], []), column_mapping=MAPPING, resolutions=[], row_decisions=[])
            self.assertEqual('formula_cache_missing', result['rows'][0]['issues'][0]['code'])
            accepted = interpret(source, catalog_context(set(OWN), [OPPONENT], []), column_mapping=MAPPING,
                                 resolutions=[resolution('attempts', '=1+1', value=2)], row_decisions=[])
            self.assertEqual('date', accepted['rows'][0]['match']['occurred_precision'])
            self.assertEqual(hashlib.sha256(before).hexdigest(), source.fingerprint)
            self.assertEqual(before, path.read_bytes())

    def test_actual_utf8_bom_csv_quote_and_explicit_mapping(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'synthetic.csv'
            with path.open('w', encoding='utf-8-sig', newline='') as stream:
                writer = csv.writer(stream); writer.writerow(FIELDS); data = values(notes='원문, 줄\n둘')
                writer.writerow([data[f] for f in FIELDS])
            before = path.read_bytes(); source = read_source(path)
            result = interpret(source, catalog_context(set(OWN), [OPPONENT], []), column_mapping=MAPPING, resolutions=[], row_decisions=[])
            self.assertEqual('원문, 줄\n둘', result['rows'][0]['match']['notes'])
            self.assertEqual(before, path.read_bytes())


if __name__ == '__main__':
    unittest.main()
